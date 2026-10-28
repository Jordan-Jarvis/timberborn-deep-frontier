"""Read-only comparison of research/dependency edits and preserved vanilla data."""
import json
from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
before=load_workbook(root/'outputs/research-workbook-before.xlsx',data_only=False)
after=load_workbook(root/'outputs/Timberborn_Deep_Frontier_Balance_Workbook.xlsx',data_only=False)
assert before.sheetnames==after.sheetnames
preserved=[]
for name in before.sheetnames:
    if not name.startswith('Vanilla'):continue
    assert list(before[name].values)==list(after[name].values),name
    preserved.append(name)
sheet=after['Buildings'];headers=[c.value for c in sheet[1]]
rows={r[0]:dict(zip(headers,r)) for r in sheet.iter_rows(min_row=2,values_only=True)}
lab=rows['B043'];catalog=json.loads((root/'tools/research_content.json').read_text());b=catalog['buildings'][0]
assert [lab[h] for h in ('Width','Length','Height')]==b['size']
assert lab['Science unlock']==12000 and lab['Power demand (hp)']==600 and lab['Power output (hp)']==0 and lab['Workers']==4
assert lab['Recipe names']=='Materials research'
recipe=[r for r in after['Recipes'].iter_rows(values_only=True) if r[0]=='MaterialsResearch'];assert len(recipe)==1
assert recipe[0][4]=='4 Gear; 2 Concrete; 1 MetalBlock; 1 Water' and recipe[0][5]=='800 Science' and recipe[0][6]==2
dep=[r for r in after['Dependencies'].iter_rows(values_only=True) if r[0]=='D017'];assert len(dep)==1 and 'Redic.NewStyles >= 0.1' in dep[0][2]
progress=[r for r in after['Progression'].iter_rows(values_only=True) if r[0]=='MaterialsResearch'];assert len(progress)==1 and progress[0][4]==12000
errors=[(s.title,c.coordinate,c.value) for s in after for row in s for c in row if c.data_type=='e'];assert not errors,errors
result={'lab':{k:lab[k] for k in ('Width','Length','Height','Workers','Science unlock','Power demand (hp)','Power output (hp)')},
        'recipe':recipe[0][:11],'dependency':dep[0][:4],'unchangedVanillaTabs':preserved,'formulaErrors':errors}
(root/'outputs/research-workbook-verification.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result,indent=2))

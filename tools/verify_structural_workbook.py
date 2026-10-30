"""Compare the structural catalog and advanced foundations with workbook data."""
import json
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
catalog = json.loads((root / "tools/structural_content.json").read_text(encoding="utf-8"))
book = load_workbook(root / "outputs/Timberborn_Deep_Frontier_Balance_Workbook.xlsx",
                     data_only=False)
sheet = book["Buildings"]
headers = [cell.value for cell in sheet[1]]
rows = {row[0]: dict(zip(headers, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)}
report = {}
for item in catalog["buildings"]:
    ident = item["id"]
    blueprint = json.loads((root / f"src/DeepFrontier/Data/Buildings/Platforms/{ident}/{ident}.IronTeeth.blueprint.json").read_text(encoding="utf-8"))
    row = rows[ident]
    side = item["kind"] == "sidedeck"
    # Side platforms report their n x n deck, one layer high; the block grid
    # adds the anchor row, the cleat layer and, for even n, an empty column
    # that keeps the anchor column flip-invariant ([n or n+1, n+1, 2]).
    shown = [item["size"][1] - 1, item["size"][1] - 1, 1] if side else item["size"]
    assert [row[h] for h in ("Width", "Length", "Height")] == shown, ident
    assert row["Science unlock"] == blueprint["BuildingSpec"]["ScienceCost"], ident
    assert row["Workers"] == row["Power demand (hp)"] == row["Power output (hp)"] == 0, ident
    assert row["Power role"] == "Passive / transport", ident
    assert "no artificial layer cap" in row["Stacking rule"], ident
    assert row["Status"] == row["Data status"] == catalog["status"], ident
    # Stackable tops: none on a bridge (a path span), only the cap of a column.
    support = 0 if item["kind"] == "bridge" else 1 if item["kind"] == "column" else shown[0] * shown[1]
    assert sum(b["Stackable"] == "BlockObject" for b in blueprint["BlockObjectSpec"]["Blocks"]) == support, ident
    # Authored costs are multiplied by their science tier's costFactor (tools/lategame_costs.mjs scaledCost).
    tiers = json.loads((root / "tools/building_construction_costs.json").read_text(encoding="utf-8"))["lategameRule"]["tiers"]
    factor = next(t["costFactor"] for t in tiers.values() if t["scienceMax"] is None or item["science"] <= t["scienceMax"])
    assert blueprint["BuildingSpec"]["BuildingCost"] == [{"Id": g, "Amount": a * factor} for g, a in item["cost"]], ident
    if item["kind"] == "scaffold":
        # Concrete platforms: k = n² + H; cost k × (3 Concrete, 2 Metal Block, 1 Gear), science 500k.
        n, depth, layers = item["size"]
        legs = layers - 1
        k = n * n + legs
        assert n == depth and n in (3, 4, 5) and legs in (1, 2, 3), ident
        assert item["cost"] == [["Concrete", 3 * k], ["MetalBlock", 2 * k], ["Gear", k]], ident
        assert item["science"] == 500 * k, ident
        assert row["Build materials"] == f"{3 * k} Concrete; {2 * k} Metal Block; {k} Gear", ident
        assert row["Placement"] == "Four supported corner columns", ident
        assert f"k = n²+H = {k}" in row["Notes"], ident
    if side:
        # Concrete side platforms: no leg layers, k = n²; same 3:2:1 bill and 500k science.
        n = item["size"][1] - 1
        k = n * n
        assert item["size"] == [n if n % 2 else n + 1, n + 1, 2] and n in (3, 4, 5), ident
        assert item["cost"] == [["Concrete", 3 * k], ["MetalBlock", 2 * k], ["Gear", k]], ident
        assert item["science"] == 500 * k, ident
        assert row["Build materials"] == f"{3 * k} Concrete; {2 * k} Metal Block; {k} Gear", ident
        assert row["Placement"].startswith("Hangs beside the top block of a reinforced dam column"), ident
        assert f"k = n² = {k}" in row["Notes"], ident
        assert blueprint["BlockObjectSpec"]["BaseZ"] == 1 and "AnchoredDeckSpec" in blueprint, ident
    report[ident] = {"size": item["size"], "science": row["Science unlock"],
                     "materials": row["Build materials"], "power": 0}
sides = sorted(item["id"] for item in catalog["buildings"] if item["kind"] == "sidedeck")
assert sides == [f"ConcreteSidePlatform{n}x{n}" for n in (3, 4, 5)], sides
scaffolds = sorted(item["id"] for item in catalog["buildings"] if item["kind"] == "scaffold")
assert scaffolds == sorted(["ConcreteScaffold3x3"] + [f"ConcreteScaffold{n}x{n}H{h}" for n in (3, 4, 5)
                                                     for h in (1, 2, 3) if (n, h) != (3, 3)]), scaffolds

advanced = {
    "B012": "Restaurant",
    "B021": "ChemicalWorks",
    "B022": "PrecisionMachineWorks",
    "B023": "WaterTreatmentPlant",
    "B026": "DeepDrill",
    "B031": "NuclearReactor",
    "B032": "UraniumProcessingPlant",
    "B033": "ElectrolysisPlant",
    "B034": "CryogenicLiquefactionPlant",
    "B036": "RocketAssemblyHangar",
    "B037": "LaunchPad",
    "B038": "MissionControl",
    "B040": "TechnicalSchool",
    "B041": "RadioactiveWasteRefinery",
    "B042": "DeepMine",
}
# Locate each exact template in authored data; duplicate/missing IDs are errors.
data = root / "src/DeepFrontier/Data/Buildings"
foundation_report = {}
for row_id, ident in advanced.items():
    paths = list(data.rglob(ident + ".IronTeeth.blueprint.json"))
    assert len(paths) == 1, (ident, paths)
    blueprint = json.loads(paths[0].read_text(encoding="utf-8"))
    concrete = [entry["Amount"] for entry in blueprint["BuildingSpec"]["BuildingCost"]
                if entry["Id"] == "Concrete"]
    assert len(concrete) == 1 and concrete[0] > 0, ident
    assert f"{concrete[0]} Concrete" in rows[row_id]["Build materials"], ident
    foundation_report[ident] = concrete[0]

for ident in ("SandstonePit", "LimestonePit", "StoneCrusher", "CementKiln", "ConcreteMixer"):
    paths = list(data.rglob(ident + ".IronTeeth.blueprint.json"))
    assert len(paths) == 1, ident
    assert all(entry["Id"] != "Concrete" for entry in
               json.loads(paths[0].read_text(encoding="utf-8"))["BuildingSpec"]["BuildingCost"]), ident

turbine = rows["ReinforcedHydroTurbine"]
hydro = json.loads((root / "tools/hydro_content.json").read_text(encoding="utf-8"))
assert turbine["Power demand (hp)"] == 0
assert turbine["Power output (hp)"] == hydro["ratedHp"]
assert turbine["Availability"] == "Authored build menu"
assert turbine["Status"] == hydro["status"]
assert [turbine[h] for h in ("Width", "Length", "Height")] == hydro["size"]
assert turbine["Science unlock"] == hydro["science"]
assert "maximum, not constant generation" in turbine["Notes"]
hydro_blueprint=json.loads((data / "Power/ReinforcedHydroTurbine/ReinforcedHydroTurbine.IronTeeth.blueprint.json").read_text(encoding="utf-8"))
assert hydro_blueprint["MechanicalNodeSpec"]["PowerOutput"]==turbine["Power output (hp)"]
support_count=sum(b["Stackable"]=="BlockObject" for b in hydro_blueprint["BlockObjectSpec"]["Blocks"])
assert turbine["Stackable"]==f"Yes; {support_count} support cells"
assert len(hydro_blueprint["HydroTurbineSpec"]["Gates"])==2
result = {"structures": report, "advancedConcreteFoundations": foundation_report,
          "starterChainNeedsNoConcrete": True,
          "turbineMatchesAuthoredMaximumWithProvisionalBalance": True}
(root / "outputs/structural-workbook-verification.json").write_text(json.dumps(result, indent=2) + "\n")
print(json.dumps(result, indent=2))

"""Read-only verification of the restaurant revision and preserved vanilla sheets."""
import json
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
before=load_workbook(root/'outputs/restaurant-workbook-baseline.xlsx',data_only=False)
after=load_workbook(root/'outputs/Timberborn_Deep_Frontier_Balance_Workbook.xlsx',data_only=False)
report={}
for name in ('VanillaAnchors','VanillaBuildings','VanillaItems','VanillaRecipes'):
    a,b=before[name],after[name]
    assert (a.max_row,a.max_column)==(b.max_row,b.max_column),name+' shape changed'
    for rows in zip(a.iter_rows(),b.iter_rows()):
        for x,y in zip(*rows):
            assert x.value==y.value,(name,x.coordinate,'value/formula changed')
            for attr in ('font','fill','border','alignment','protection','number_format'):
                assert copy(getattr(x,attr))==copy(getattr(y,attr)),(name,x.coordinate,attr+' changed')
    assert set(a.tables)==set(b.tables),name+' table names changed'
    for table in a.tables:
        assert str(a.tables[table])==str(b.tables[table]),name+' table changed'
    report[name]={'valuesAndFormulasUnchanged':True,'cellFormatsUnchanged':True,'tablesUnchanged':True}
buildings=after['Buildings']
row=next(r for r in buildings.iter_rows(values_only=True) if r[1]=='Restaurant')
columns={c.value:i for i,c in enumerate(buildings[1])}
assert row[columns['Power demand (hp)']]==400
assert row[columns['Power output (hp)']]==0
assert row[columns['Workers']]==4
assert row[columns['Science unlock']]==800
recipes=[r for r in after['Recipes'].iter_rows(values_only=True) if r[1]=='Restaurant']
assert len(recipes)==5
assert {r[0] for r in recipes}=={'R016','FermentedCassava','FermentedSoybean','FermentedMushroom','SweetBerryCake'}
cake=next(r for r in after['Items'].iter_rows(values_only=True) if r[1]=='Sweet Berry Cake')
assert cake[5]=='RestaurantDish' and cake[7]=='Restaurant only'
errors=[]
for sheet in after:
    for row in sheet:
        for cell in row:
            if cell.data_type=='e':errors.append([sheet.title,cell.coordinate,cell.value])
assert not errors,errors
report.update(restaurantPowerDemandHp=400,restaurantPowerOutputHp=0,restaurantRecipes=len(recipes),cakeStorage=cake[7],formulaErrors=errors)
(root/'outputs/restaurant-workbook-verification.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps(report,indent=2))
